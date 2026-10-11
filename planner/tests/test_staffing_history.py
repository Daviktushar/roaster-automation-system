from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings

from openpyxl import Workbook, load_workbook
from planner.forms import StaffingConfigurationForm
from planner.models import Employee, EmployeePatternChange, GenerationRun, HistoricalRosterSnapshot, PatternHistoryBaseline, RosterSheet, StaffingRequirement, WorkbookImport
from planner.services.history import history_status, import_workbook_history, would_exceed_annual_pattern_limit
from planner.services.roster_generation import _build_sheet_assignments
from planner.services.roster_generation import create_roster_preview
from django.urls import reverse


class StaffingAndHistoryTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="owner", password="test-pass")

    def test_staffing_configuration_is_persisted_and_reused(self):
        form = StaffingConfigurationForm(data={
            "effective_month": "2026-11",
            "requirements": "HDC3,2,weekday,4\nHDC3,3,everyday,1",
            "confirmation_note": "Confirmed by roster owner from reviewed requirements",
        }, owner=self.user)
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        self.assertEqual(StaffingRequirement.objects.filter(owner=self.user).count(), 3)
        again = StaffingConfigurationForm(owner=self.user)
        self.assertIn("HDC3,2,weekday,4", again["requirements"].value())

    def test_specific_date_requirement_is_saved_as_override(self):
        form = StaffingConfigurationForm(data={
            "effective_month": "2026-11",
            "requirements": "HDC3,2,weekday,4\nHDC3,2,2026-11-02,3",
            "confirmation_note": "Reviewed date-specific change",
        }, owner=self.user)
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        override = StaffingRequirement.objects.get(owner=self.user, specific_date=date(2026, 11, 2))
        self.assertEqual(override.required_count, 3)

    def test_staffing_shortage_is_actionable(self):
        rows = [{
            "employee": "Engineer A", "activity": "Incidents", "source_row": 3,
            "assignments": {str(day): "G" for day in range(1, 32)},
            "source_assignments": {str(day): "G" for day in range(1, 32)},
            "previous_shifts": [], "rotation_target": None, "profile": None,
            "excluded": False, "previous_shift": None,
        }]
        output, warnings = _build_sheet_assignments("HDC3", rows, date(2026, 10, 1), date(2026, 11, 1), set(), {("HDC3", "2", "weekday"): 2})
        shortage = next(item for item in warnings if item["code"] == "staffing_shortage" and item["required"] == 2)
        self.assertEqual(shortage["required"], 2)
        self.assertLess(shortage["actual"], 2)

    def test_annual_pattern_change_limit_counts_persisted_changes(self):
        employee = Employee.objects.create(owner=self.user, display_name="Pattern Agent", normalized_name="pattern agent")
        for month, previous, new in ((2, "1", "2"), (5, "2", "3")):
            EmployeePatternChange.objects.create(employee=employee, effective_start=date(2026, month, 1), previous_pattern=previous, new_pattern=new, changed_on=date(2026, month, 1), source="test")
        exceeds, count = would_exceed_annual_pattern_limit(employee, "3", "1", date(2026, 11, 1))
        self.assertTrue(exceeds)
        self.assertEqual(count, 2)

    def test_workbook_history_import_and_baseline_completeness(self):
        workbook = WorkbookImport.objects.create(uploaded_by=self.user, source_file="workbooks/old.xlsx", source_filename="old.xlsx")
        employee = Employee.objects.create(owner=self.user, display_name="Alex", normalized_name="alex")
        sheet = RosterSheet.objects.create(workbook=workbook, name="HDC3", roster_month=date(2026, 10, 1), layout_kind="roster", data_rows=[{
            "name": "Alex", "category": "Incidents", "transport_preference": "OWN", "assignments": {str(day): "3" for day in range(1, 32)},
        }])
        self.assertEqual(import_workbook_history(workbook), 1)
        snapshot = HistoricalRosterSnapshot.objects.get(roster_sheet=sheet, employee=employee)
        self.assertEqual(snapshot.pattern, "3")
        self.assertFalse(history_status(self.user, date(2026, 11, 1))["complete_for_target_year"])

    def test_staffing_and_baseline_screens_are_available(self):
        workbook = WorkbookImport.objects.create(uploaded_by=self.user, source_file="workbooks/old.xlsx", source_filename="old.xlsx")
        self.client.force_login(self.user)
        response = self.client.get(reverse("configure_staffing", args=[workbook.pk]))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(reverse("configure_staffing", args=[workbook.pk]), {
            "effective_month": "2026-11", "requirements": "HDC3,2,weekday,2", "confirmation_note": "Owner confirmed",
        })
        self.assertEqual(response.status_code, 302)
        self.assertTrue(StaffingRequirement.objects.filter(owner=self.user, facility="HDC3", shift_code="2").exists())
        response = self.client.get(reverse("confirm_pattern_baseline", args=[workbook.pk]))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(reverse("confirm_pattern_baseline", args=[workbook.pk]), {"effective_from": "2026-01-01", "note": "Earlier records unavailable; start prospective tracking."})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(PatternHistoryBaseline.objects.filter(owner=self.user, effective_from=date(2026, 1, 1)).exists())
        self.assertFalse(history_status(self.user, date(2026, 11, 1))["complete_for_target_year"])

    def test_next_month_generation_uses_saved_staffing_and_baseline(self):
        workbook = WorkbookImport.objects.create(uploaded_by=self.user, source_file="workbooks/dec.xlsx", source_filename="dec.xlsx")
        rows = []
        for name, last_off in (("Alex", False), ("Blair", True)):
            assignments = {str(day): ("WO" if date(2026, 12, day).weekday() >= 5 or (last_off and day == 31) else "G") for day in range(1, 32)}
            rows.append({"name": name, "category": "HDC5", "source_row": len(rows) + 3, "assignments": assignments, "transport_preference": "OWN" if name == "Alex" else "CAB"})
            Employee.objects.create(owner=self.user, display_name=name, normalized_name=name.casefold(), transport_preference=rows[-1]["transport_preference"])
        RosterSheet.objects.create(workbook=workbook, name="HDC5", roster_month=date(2026, 12, 1), layout_kind="roster", data_rows=rows)
        PatternHistoryBaseline.objects.create(owner=self.user, effective_from=date(2026, 1, 1), note="Prospective change tracking; earlier history not asserted.", confirmed_by=self.user)
        target = date(2027, 1, 1)
        StaffingRequirement.objects.create(owner=self.user, facility="HDC5", shift_code="G", day_type="weekday", required_count=1, effective_from=target, confirmation_note="Confirmed")
        self.assertTrue(history_status(self.user, target)["complete_for_target_year"])
        result, warnings, _, _ = create_roster_preview(workbook, target, ["HDC5"], staffing_requirements={("HDC5", "G", "weekday"): 1}, history_complete=True)
        self.assertEqual(len(result), 2)
        self.assertFalse([item for item in warnings if item.get("severity") == "error"])

    def test_configured_generation_downloads_without_manual_approval(self):
        storage_settings = {"default": {"BACKEND": "planner.tests.storage.MemoryStorage"}, "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"}}
        with override_settings(STORAGES=storage_settings):
            self.client.force_login(self.user)
            book = Workbook()
            sheet = book.active
            sheet.title = "HDC5"
            sheet["A1"] = "Dec'26"
            sheet["B1"] = "HDC5"
            sheet["A2"] = "EMPLOYEE NAME"
            sheet["B2"] = "Facility"
            for day in range(1, 32):
                current = date(2026, 12, day)
                sheet.cell(1, day + 2, day)
                sheet.cell(2, day + 2, current.strftime("%a"))
                for row, name in ((3, "Alex"), (4, "Blair")):
                    sheet.cell(row, 1, name)
                    sheet.cell(row, 2, "HDC5")
                    sheet.cell(row, day + 2, "WO" if current.weekday() >= 5 or (name == "Blair" and day == 31) else "G")
            source = BytesIO()
            book.save(source)
            response = self.client.post("/workbooks/upload/", {"workbook": SimpleUploadedFile("december.xlsx", source.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            self.assertEqual(response.status_code, 302)
            imported = WorkbookImport.objects.get(source_filename="december.xlsx")
            response = self.client.post(reverse("configure_staffing", args=[imported.pk]), {"effective_month": "2027-01", "requirements": "HDC5,G,weekday,1", "confirmation_note": "Confirmed by roster owner"})
            self.assertEqual(response.status_code, 302)
            response = self.client.post(reverse("confirm_pattern_baseline", args=[imported.pk]), {"effective_from": "2027-01-01", "note": "Prospective tracking; earlier records are not asserted."})
            self.assertEqual(response.status_code, 302)
            response = self.client.post("/rosters/generate/", {"source_id": imported.pk, "target_month": "1", "target_year": "2027", "location": "HDC5"})
            self.assertEqual(response.status_code, 302)
            run = GenerationRun.objects.filter(source=imported).order_by("-created_at").first()
            self.assertEqual(run.validation_status, "NEEDS_ATTENTION")
            self.assertEqual(run.status, "review")
            response = self.client.get(reverse("download_run", args=[run.pk]))
            self.assertEqual(response.status_code, 200)
            exported = load_workbook(BytesIO(b"".join(response.streaming_content)), read_only=True, data_only=True)
            self.assertEqual(exported["HDC5"]["A1"].value, "January 2027")
            self.assertEqual(exported["Validation Summary"]["B3"].value, "NEEDS_ATTENTION")
            self.assertEqual(exported["HDC5"]["A3"].value, "Alex")
            self.assertTrue(any(cell.value == "WO" for row in exported["HDC5"].iter_rows(min_row=3, min_col=3) for cell in row))
            exported.close()

    def test_previous_roster_generates_downloadable_provisional_without_manual_setup(self):
        storage_settings = {"default": {"BACKEND": "planner.tests.storage.MemoryStorage"}, "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"}}
        with override_settings(STORAGES=storage_settings):
            self.client.force_login(self.user)
            book = Workbook()
            sheet = book.active
            sheet.title = "HDC5"
            sheet["A1"] = "Dec'26"
            sheet["B1"] = "HDC5"
            sheet["A2"] = "EMPLOYEE NAME"
            sheet["B2"] = "Facility"
            for day in range(1, 32):
                current = date(2026, 12, day)
                sheet.cell(1, day + 2, day)
                sheet.cell(2, day + 2, current.strftime("%a"))
                for row, name in ((3, "Alex"), (4, "Blair")):
                    sheet.cell(row, 1, name)
                    sheet.cell(row, 2, "HDC5")
                    sheet.cell(row, day + 2, "WO" if current.weekday() >= 5 or (name == "Blair" and day == 31) else "G")
            source = BytesIO()
            book.save(source)
            response = self.client.post("/workbooks/upload/", {"workbook": SimpleUploadedFile("previous.xlsx", source.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            self.assertEqual(response.status_code, 302)
            imported = WorkbookImport.objects.get(source_filename="previous.xlsx")
            run = GenerationRun.objects.get(source=imported)
            self.assertEqual(run.validation_status, "NEEDS_ATTENTION")
            self.assertEqual(run.validation_report["staffing"], "INFERRED")
            self.assertTrue(any(item["code"] == "annual_history_unavailable" for item in run.warnings))

            response = self.client.get(reverse("run_detail", args=[run.pk]))
            self.assertContains(response, "Download provisional Excel")
            response = self.client.get(reverse("download_run", args=[run.pk]))
            self.assertEqual(response.status_code, 200)
            exported = load_workbook(BytesIO(b"".join(response.streaming_content)), read_only=True, data_only=True)
            self.assertEqual(exported["HDC5"]["A1"].value, "January 2027")
            self.assertEqual(exported["Validation Summary"]["B3"].value, "NEEDS_ATTENTION")
            self.assertEqual(exported["HDC5"]["A3"].value, "Alex")
            exported.close()
