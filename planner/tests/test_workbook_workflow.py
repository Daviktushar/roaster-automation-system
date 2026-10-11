from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook

from planner.models import Employee, EmployeeAssignment, GenerationRun, WorkbookImport


def make_reference_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "HDC3"
    sheet["A1"] = "Oct'26"
    sheet["B1"] = "HDC3"
    sheet["A2"] = "EMPLOYEE NAME"
    sheet["B2"] = "Activity"
    for day in range(1, 32):
        col = day + 2
        current = date(2026, 10, day)
        sheet.cell(1, col, day)
        sheet.cell(2, col, current.strftime("%a"))
        sheet.cell(3, col, "WO" if current.weekday() >= 5 else "G")
    sheet["A3"] = "Demo Agent"
    sheet["B3"] = "Imaging"
    sheet["A4"] = "Leave Agent"
    sheet["B4"] = "Imaging"
    for day in range(1, 32):
        current = date(2026, 10, day)
        sheet.cell(4, day + 2, "WO" if current.weekday() >= 5 else "G")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def make_consolidated_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Roster Data"
    sheet["A1"] = "Oct'26"
    sheet["A2"] = "Location"
    sheet["B2"] = "Employee Name"
    sheet["C2"] = "Activity"
    employees = ((3, "HDC3", "HDC3 Agent", "Operations"), (4, "HDC4", "HDC4 Agent", "Operations"))
    for day in range(1, 32):
        column = day + 3
        current = date(2026, 10, day)
        sheet.cell(1, column, day)
        sheet.cell(2, column, current.strftime("%a"))
        for row_number, location, name, activity in employees:
            sheet.cell(row_number, 1, location)
            sheet.cell(row_number, 2, name)
            sheet.cell(row_number, 3, activity)
            sheet.cell(row_number, column, "WO" if current.weekday() >= 5 else "G")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class WorkbookWorkflowTests(TestCase):
    def test_invalid_upload_shows_inline_error_and_keeps_dashboard(self):
        user = get_user_model().objects.create_user(username="invalid-upload", password="Roster-Preview-482!")
        self.client.force_login(user)
        response = self.client.post("/workbooks/upload/", {
            "workbook": SimpleUploadedFile("invalid.csv", b"not an Excel workbook", content_type="text/csv"),
        })
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Choose an .xlsx or .xlsm workbook.", status_code=400)
        self.assertEqual(WorkbookImport.objects.count(), 0)
        response = self.client.post("/workbooks/upload/", {
            "workbook": SimpleUploadedFile("broken.xlsx", b"not an Excel package", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        })
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Excel could not read this workbook", status_code=400)
        self.assertEqual(WorkbookImport.objects.count(), 0)

    def test_consolidated_upload_generates_selected_location_tab(self):
        storage_settings = {
            "default": {"BACKEND": "planner.tests.storage.MemoryStorage"},
            "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
        }
        with override_settings(STORAGES=storage_settings):
            user = get_user_model().objects.create_user(username="consolidated", password="Roster-Preview-482!")
            self.client.force_login(user)
            response = self.client.post("/workbooks/upload/", {
                "workbook": SimpleUploadedFile("all_locations.xlsx", make_consolidated_workbook(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            })
            self.assertEqual(response.status_code, 302)
            source = WorkbookImport.objects.get()
            self.assertEqual(set(source.sheets.filter(layout_kind="roster").values_list("name", flat=True)), {"HDC3", "HDC4"})
            page = self.client.get(f"/workbooks/{source.pk}/")
            self.assertContains(page, '<option value="HDC3"')
            self.assertContains(page, '<option value="HDC4"')

            response = self.client.post("/rosters/generate/", {
                "source_id": source.pk, "location": "HDC3", "target_month": "11", "target_year": "2026",
            })
            self.assertEqual(response.status_code, 302)
            run = GenerationRun.objects.get(source=source)
            self.assertEqual(run.categories, ["HDC3"])
            with run.output_file.open("rb") as generated:
                output = load_workbook(generated, read_only=True, data_only=True)
                self.assertIn("HDC3", output.sheetnames)
                self.assertNotIn("HDC4", output.sheetnames)
                self.assertEqual(output["HDC3"]["A1"].value, "HDC3 roster · November 2026")
                output.close()

    def test_upload_project_and_download_a_separate_workbook(self):
        storage_settings = {
            "default": {"BACKEND": "planner.tests.storage.MemoryStorage"},
            "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
        }
        with override_settings(STORAGES=storage_settings):
            user = get_user_model().objects.create_user(username="planner", password="Roster-Preview-482!")
            self.client.force_login(user)
            source_bytes = make_reference_workbook()
            response = self.client.post("/workbooks/upload/", {
                "workbook": SimpleUploadedFile("sample.xlsx", source_bytes, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            })
            self.assertEqual(response.status_code, 302)
            source = WorkbookImport.objects.get()
            self.assertEqual(source.sheets.get().name, "HDC3")
            self.assertTrue(source.source_file.storage.exists(source.source_file.name))
            employee = Employee.objects.get(display_name="Demo Agent")
            self.assertEqual(employee.owner, user)
            generation_page = self.client.get(f"/workbooks/{source.pk}/")
            self.assertContains(generation_page, 'id="id_location"')
            self.assertContains(generation_page, '<option value="HDC3"')
            self.assertContains(generation_page, 'id="id_target_month"')
            self.assertContains(generation_page, 'id="id_target_year"')

            response = self.client.get(f"/workbooks/{source.pk}/employees/{employee.pk}/edit/")
            self.assertEqual(response.status_code, 200)
            response = self.client.post(f"/workbooks/{source.pk}/employees/{employee.pk}/edit/", {
                "display_name": "Alex Doe",
                "email": "alex@example.com",
                "phone": "555-0123",
                "notes": "Updated before roster generation.",
                "active": "on",
            })
            self.assertEqual(response.status_code, 302)
            employee.refresh_from_db()
            self.assertEqual(employee.display_name, "Alex Doe")
            self.assertTrue(employee.active)

            leave_employee = Employee.objects.get(display_name="Leave Agent")
            response = self.client.post(f"/workbooks/{source.pk}/employees/{leave_employee.pk}/edit/", {
                "display_name": "Leave Agent",
                "email": "",
                "phone": "",
                "notes": "",
                "active": "on",
                "planned_leave": "2026-11-05 to 2026-11-06",
            })
            self.assertEqual(response.status_code, 302)
            leave_employee.refresh_from_db()
            self.assertEqual(leave_employee.planned_leave, [{"start": "2026-11-05", "end": "2026-11-06"}])

            response = self.client.post(f"/workbooks/{source.pk}/employees/{employee.pk}/roster/", {"action": "remove"})
            self.assertEqual(response.status_code, 302)
            self.assertTrue(EmployeeAssignment.objects.get(employee=employee).excluded_from_roster)
            self.assertContains(self.client.get(f"/workbooks/{source.pk}/"), "Removed")

            response = self.client.post("/rosters/generate/", {
                "source_id": str(source.pk),
                "target_month": "11",
                "target_year": "2026",
                "public_holidays": "2026-11-04",
                "location": "HDC3",
            })
            self.assertEqual(response.status_code, 302)
            run = GenerationRun.objects.get()
            self.assertEqual(run.target_month, date(2026, 11, 1))
            self.assertEqual(run.status, "review")
            self.assertEqual(run.assignments[0]["employee"], "Alex Doe")
            self.assertTrue(run.assignments[0]["active"])
            self.assertFalse(run.assignments[0]["included"])
            self.assertTrue(all(not code for code in run.assignments[0]["assignments"]))
            # Leave on Nov 5-6 blocks the adjacent Nov 4 holiday and Nov 7-8 weekend.
            self.assertFalse(any(code in {"G", "1", "2", "3"} for code in run.assignments[1]["assignments"][3:8]))
            self.assertFalse(any(item["code"] == "holiday_intake_out_of_scope" for item in run.warnings))
            self.assertTrue(run.output_file.storage.exists(run.output_file.name))

            response = self.client.get(f"/rosters/{run.pk}/download/")
            self.assertEqual(response.status_code, 302)
            self.assertTrue(any(item.get("severity") == "error" for item in run.warnings))

            response = self.client.post(f"/workbooks/{source.pk}/employees/{employee.pk}/roster/", {"action": "restore"})
            self.assertEqual(response.status_code, 302)
            self.assertFalse(EmployeeAssignment.objects.get(employee=employee).excluded_from_roster)
            with source.source_file.open("rb") as original:
                source_book = load_workbook(original, read_only=True, data_only=False)
                self.assertEqual(source_book["HDC3"]["A1"].value, "Oct'26")
                source_book.close()
