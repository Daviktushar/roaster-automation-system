from datetime import date
from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from planner.services.excel_export import export_workbook


class ExcelExportTests(unittest.TestCase):
    def test_export_patches_copy_and_keeps_existing_style_and_formula(self):
        source = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "HDC3"
        sheet["A1"] = "Oct'26"
        sheet["B1"] = "HDC3"
        sheet["A2"] = "EMPLOYEE NAME"
        sheet["B2"] = "Activity"
        for day in range(1, 32):
            col = day + 2
            sheet.cell(1, col, day)
            sheet.cell(2, col, "Mon")
            cell = sheet.cell(3, col, "G")
            cell.fill = PatternFill("solid", fgColor="C6EFCE")
        sheet["AH3"] = "=COUNTIF(C3:AG3,\"G\")"
        workbook.save(source)
        source.seek(0)
        row = {"sheet": "HDC3", "source_row": 3, "assignments": ["WO"] + ["G"] * 29}
        result, _ = export_workbook(source, [{"name": "HDC3", "day_count": 31}], [row], date(2026, 11, 1))
        exported = load_workbook(result, data_only=False)
        output_sheet = exported["HDC3"]
        self.assertEqual(output_sheet["A1"].value, "November 2026")
        self.assertEqual(output_sheet["C3"].value, "WO")
        self.assertEqual(output_sheet["C3"].fill.fgColor.rgb[-6:], "C6EFCE")
        self.assertEqual(output_sheet["AH3"].value, '=COUNTIF(C3:AG3,"G")')
        exported.close()

    def test_export_adds_validation_and_coverage_summary_sheet(self):
        source = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "HDC3"
        sheet["A1"] = "Oct'26"
        sheet["A2"] = "EMPLOYEE NAME"
        for day in range(1, 32):
            sheet.cell(1, day + 2, day)
            sheet.cell(2, day + 2, "Mon")
        sheet["A3"] = "Alex"
        workbook.save(source)
        source.seek(0)
        report = {"overall": "NEEDS_ATTENTION", "rules": "VALID", "staffing": "INCOMPLETE", "history": "INCOMPLETE", "feasibility": "FEASIBLE", "coverage": [{"facility": "HDC3", "date": "2026-11-02", "shift": "2", "required": 2, "actual": 1, "shortage": 1, "excess": 0}]}
        result, _ = export_workbook(source, [{"name": "HDC3", "day_count": 31}], [], date(2026, 11, 1), report, [{"severity": "warning", "code": "incomplete", "message": "Setup incomplete"}])
        exported = load_workbook(result, data_only=True)
        summary = exported["Validation Summary"]
        self.assertEqual(summary["B3"].value, "NEEDS_ATTENTION")
        self.assertEqual(summary["D10"].value, 2)
        self.assertEqual(summary["E10"].value, 1)
        self.assertEqual(exported["HDC3"]["A1"].value, "November 2026")
        exported.close()


if __name__ == "__main__":
    unittest.main()
