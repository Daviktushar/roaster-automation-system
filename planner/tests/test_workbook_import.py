from datetime import date
from io import BytesIO
import unittest

from openpyxl import Workbook

from planner.services.workbook_import import parse_workbook


class NamedUpload(BytesIO):
    name = "sample.xlsx"


class WorkbookImportTests(unittest.TestCase):
    def test_month_day_codes_and_exact_trailing_space_are_detected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "DD "
        sheet["A1"] = "Oct'26"
        sheet["B1"] = "DD"
        sheet["A2"] = "EMPLOYEE NAME"
        sheet["B2"] = "Activity"
        for day in range(1, 32):
            sheet.cell(1, day + 2, day)
            sheet.cell(2, day + 2, "Thu")
        sheet["A3"] = "Demo employee"
        sheet["B3"] = "META"
        sheet["C3"] = "G"
        sheet["D3"] = "WO"
        source = NamedUpload()
        workbook.save(source)
        source.seek(0)
        result = parse_workbook(source)
        self.assertEqual(result["sheet_summary"][0]["name"], "DD ")
        self.assertEqual(result["sheet_summary"][0]["month"], date(2026, 10, 1).isoformat())
        self.assertEqual(result["sheets"][0]["rows"][0]["assignments"]["1"], "G")
        self.assertEqual(result["sheets"][0]["rows"][0]["source_row"], 3)

    def test_non_excel_extension_is_rejected(self):
        file = NamedUpload()
        file.name = "roster.csv"
        with self.assertRaises(ValueError):
            parse_workbook(file)

    def test_ambiguous_base_file_rows_are_retained_for_admin_confirmation(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Base file"
        sheet["H1"] = "2nd Shift"
        sheet["I1"] = "HDC3"
        sheet["J1"] = "HDC4"
        sheet["H2"] = 3
        sheet["I2"] = 4
        sheet["J2"] = 2
        source = NamedUpload()
        workbook.save(source)
        source.seek(0)
        result = parse_workbook(source)
        base = result["sheets"][0]
        self.assertEqual(base["layout_kind"], "staffing_reference")
        self.assertEqual(base["rows"][0]["cells"], [
            {"column": "H", "value": "2nd Shift"},
            {"column": "I", "value": "HDC3"},
            {"column": "J", "value": "HDC4"},
        ])


if __name__ == "__main__":
    unittest.main()
