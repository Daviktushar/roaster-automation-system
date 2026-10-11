"""Patch roster values into a copy of the source XLSX package.

The ZIP/XML approach avoids expanding HDC3's formatting-only rows and retains
the original sheet XML, formulas, styles, hidden state, merged cells, and widths.
"""
from __future__ import annotations

import calendar
import html
import re
import zipfile
from datetime import date
from io import BytesIO
from xml.sax.saxutils import escape
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CELL_RE = re.compile(rb"<c(?=[\s>])(?P<attrs>[^>]*?)(?:/\s*>|>(?P<body>.*?)</c>)", re.DOTALL)
REF_RE = re.compile(rb'\br="([A-Z]{1,3}[0-9]+)"')


def _sheet_paths(package):
    workbook_root = ET.fromstring(package.read("xl/workbook.xml"))
    relationships = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
    rel_map = {item.attrib["Id"]: item.attrib["Target"] for item in relationships.findall(f"{{{PKG_REL_NS}}}Relationship")}
    result = {}
    for item in workbook_root.findall(f"{{{MAIN_NS}}}sheets/{{{MAIN_NS}}}sheet"):
        relationship_id = item.attrib[f"{{{REL_NS}}}id"]
        target = rel_map[relationship_id].replace("\\", "/")
        if target.startswith("/"):
            path = target.lstrip("/")
        elif target.startswith("xl/"):
            path = target
        else:
            path = "xl/" + target
        result[item.attrib["name"]] = path
    return result


def _cell(value):
    if value is None:
        return b""
    if isinstance(value, (int, float)):
        return f"<v>{value}</v>".encode("ascii")
    text = html.escape(str(value), quote=False)
    return f'<is><t xml:space="preserve">{text}</t></is>'.encode("utf-8")


def _patch_sheet(xml_bytes, updates):
    remaining = set(updates)
    changed = 0

    def replace(match):
        nonlocal changed
        attrs = match.group("attrs")
        reference_match = REF_RE.search(attrs)
        if not reference_match:
            return match.group(0)
        reference = reference_match.group(1).decode("ascii")
        if reference not in updates:
            return match.group(0)
        remaining.discard(reference)
        changed += 1
        value = updates[reference]
        if value is None:
            # Keep the existing styled cell, clearing only the cell's content.
            attrs = re.sub(rb'\s+t="[^"]*"', b"", attrs)
            return b"<c" + attrs + b"/>"
        attrs = re.sub(rb'\s+t="[^"]*"', b"", attrs)
        if isinstance(value, str):
            attrs += b' t="inlineStr"'
        return b"<c" + attrs + b">" + _cell(value) + b"</c>"

    patched = CELL_RE.sub(replace, xml_bytes)
    return patched, changed, sorted(remaining)


def _excel_month_label(target_month):
    return f"{calendar.month_name[target_month.month]} {target_month.year}"


def _validation_sheet_xml(target_month, validation_report, warnings):
    rows = [
        ["Roster Validation Summary"],
        ["Target month", _excel_month_label(target_month)],
        ["Overall status", validation_report.get("overall", "NEEDS_ATTENTION")],
        ["Rule compliance", validation_report.get("rules", "NOT VERIFIED")],
        ["Staffing", validation_report.get("staffing", "NOT VERIFIED")],
        ["Historical completeness", validation_report.get("history", "NOT VERIFIED")],
        ["Generation feasibility", validation_report.get("feasibility", "NOT VERIFIED")],
        [],
        ["Facility", "Date", "Shift", "Required", "Assigned", "Shortage", "Excess"],
    ]
    for item in validation_report.get("coverage", []):
        rows.append([item.get("facility"), item.get("date"), item.get("shift"), item.get("required"), item.get("actual"), item.get("shortage"), item.get("excess")])
    rows.extend([[], ["Validation findings"]])
    for item in warnings:
        rows.append([item.get("severity", "info").upper(), item.get("code", ""), item.get("message", "")])

    def col_name(index):
        result = ""
        while index:
            index, remainder = divmod(index - 1, 26)
            result = chr(65 + remainder) + result
        return result

    row_xml = []
    for row_index, values in enumerate(rows, 1):
        cells = []
        for col_index, value in enumerate(values, 1):
            if value is None:
                continue
            ref = f"{col_name(col_index)}{row_index}"
            if isinstance(value, (int, float)):
                cells.append(f'<c r="{ref}"><v>{value}</v></c>')
            else:
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>')
        row_xml.append(f'<row r="{row_index}">{"".join(cells)}</row>')
    last_col = col_name(max(len(row) for row in rows))
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="{MAIN_NS}"><dimension ref="A1:{last_col}{len(rows)}"/><cols><col min="1" max="1" width="22" customWidth="1"/><col min="2" max="2" width="20" customWidth="1"/><col min="3" max="7" width="16" customWidth="1"/></cols><sheetData>{"".join(row_xml)}</sheetData></worksheet>').encode("utf-8")


def export_workbook(source_stream, source_sheets, assignment_rows, target_month, validation_report=None, warnings=None):
    """Return a new workbook with selected roster sheets projected for a month."""
    source_stream.seek(0)
    input_bytes = source_stream.read()
    output = BytesIO()
    all_missing = []
    patched_count = 0
    target_days = calendar.monthrange(target_month.year, target_month.month)[1]
    validation_report = validation_report or {}
    warnings = warnings or []
    validation_sheet_path = "xl/worksheets/sheet999.xml"
    add_validation_sheet = bool(validation_report)

    with zipfile.ZipFile(BytesIO(input_bytes), "r") as source_zip, zipfile.ZipFile(output, "w") as output_zip:
        sheet_paths = _sheet_paths(source_zip)
        updates_by_path = {}
        for sheet in source_sheets:
            sheet_name = sheet["name"]
            xml_path = sheet_paths.get(sheet_name)
            if not xml_path:
                all_missing.append(f"Worksheet {sheet_name!r} could not be mapped to workbook XML.")
                continue
            day_count = int(sheet.get("day_count", 0))
            updates = {"A1": _excel_month_label(target_month)}
            writable_days = min(day_count, 31)
            for day_number in range(1, writable_days + 1):
                date_col = get_column_letter(day_number + 2)
                weekday_col = date_col
                if day_number <= target_days:
                    current_date = date(target_month.year, target_month.month, day_number)
                    updates[f"{date_col}1"] = day_number
                    updates[f"{weekday_col}2"] = current_date.strftime("%a")
                else:
                    updates[f"{date_col}1"] = None
                    updates[f"{weekday_col}2"] = None
            for row in assignment_rows:
                if row["sheet"] != sheet_name:
                    continue
                source_row = int(row["source_row"])
                updates[f"A{source_row}"] = row.get("employee", "")
                identity_column = row.get("employee_id_column")
                if identity_column is not None and row.get("employee_code"):
                    updates[f"{get_column_letter(int(identity_column) + 1)}{source_row}"] = row["employee_code"]
                for day_number in range(1, min(target_days, writable_days) + 1):
                    column = get_column_letter(day_number + 2)
                    values = row["assignments"]
                    updates[f"{column}{source_row}"] = values[day_number - 1] if day_number <= len(values) else None
                for day_number in range(target_days + 1, writable_days + 1):
                    column = get_column_letter(day_number + 2)
                    updates[f"{column}{source_row}"] = None
            updates_by_path[xml_path] = updates

        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename in updates_by_path:
                data, count, missing = _patch_sheet(data, updates_by_path[info.filename])
                patched_count += count
                updates = updates_by_path[info.filename]
                required_missing = [cell for cell in missing if updates.get(cell) not in (None, "")]
                all_missing.extend(f"Cell {cell} was absent from {info.filename}; formatting/layout was left intact." for cell in required_missing)
            elif info.filename == "xl/workbook.xml":
                if re.search(rb"<calcPr\b[^>]*/>", data):
                    def set_recalculation(match):
                        attrs = match.group(1)
                        for attribute in (b"fullCalcOnLoad", b"forceFullCalc", b"calcMode"):
                            attrs = re.sub(rb"\s+" + attribute + rb'="[^"]*"', b"", attrs)
                        return b'<calcPr' + attrs + b' fullCalcOnLoad="1" forceFullCalc="1" calcMode="auto"/>'
                    data = re.sub(rb"<calcPr\b([^>]*)/>", set_recalculation, data, count=1)
                else:
                    data = data.replace(b"</workbook>", b'<calcPr fullCalcOnLoad="1" forceFullCalc="1" calcMode="auto"/></workbook>')
                if add_validation_sheet:
                    sheet_ids = [int(value) for value in re.findall(rb'<sheet\b[^>]*\bsheetId="(\d+)"', data)]
                    relationship_data = source_zip.read("xl/_rels/workbook.xml.rels")
                    relationship_ids = [int(value) for value in re.findall(rb'\bId="rId(\d+)"', relationship_data)]
                    sheet_id = max(sheet_ids or [0]) + 1
                    rel_id = max(relationship_ids or [0]) + 1
                    name = b'Validation Summary'
                    new_sheet = f'<sheet name="{name.decode()}" sheetId="{sheet_id}" state="visible" r:id="rId{rel_id}"/>'.encode("ascii")
                    data = data.replace(b"</sheets>", new_sheet + b"</sheets>")
            elif info.filename == "xl/_rels/workbook.xml.rels" and add_validation_sheet:
                rel_ids = [int(value) for value in re.findall(rb'\bId="rId(\d+)"', data)]
                rel_id = max(rel_ids or [0]) + 1
                sheet_ids = [int(value) for value in re.findall(rb"sheet(\d+)\.xml", b" ".join(i.filename.encode() for i in source_zip.infolist()))]
                sheet_number = max(sheet_ids or [0]) + 1
                validation_sheet_path = f"xl/worksheets/sheet{sheet_number}.xml"
                relationship = f'<Relationship Id="rId{rel_id}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{sheet_number}.xml"/>'.encode("ascii")
                data = data.replace(b"</Relationships>", relationship + b"</Relationships>")
            elif info.filename == "[Content_Types].xml" and add_validation_sheet:
                sheet_ids = [int(value) for value in re.findall(rb"sheet(\d+)\.xml", b" ".join(i.filename.encode() for i in source_zip.infolist()))]
                sheet_number = max(sheet_ids or [0]) + 1
                override = f'<Override PartName="/xl/worksheets/sheet{sheet_number}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'.encode("ascii")
                data = data.replace(b"</Types>", override + b"</Types>")
            output_zip.writestr(info, data)
        if add_validation_sheet:
            output_zip.writestr(validation_sheet_path, _validation_sheet_xml(target_month, validation_report, warnings))
    output.seek(0)
    return output, {"updated_cells": patched_count, "layout_warnings": all_missing}


def export_consolidated_workbook(assignment_rows, target_month, validation_report=None, warnings=None):
    """Create a multi-tab roster when the uploaded source combines locations."""
    validation_report = validation_report or {}
    warnings = warnings or []
    workbook = Workbook()
    workbook.remove(workbook.active)
    facilities = list(dict.fromkeys(row.get("sheet", "Location") for row in assignment_rows))
    used_titles = set()
    for facility in facilities:
        base_title = re.sub(r"[\\/*?:\[\]]", " ", str(facility)).strip()[:31] or "Location"
        safe_title = base_title
        suffix = 2
        while safe_title.casefold() in used_titles or safe_title.casefold() == "validation summary":
            tail = f" ({suffix})"
            safe_title = f"{base_title[:31 - len(tail)]}{tail}"
            suffix += 1
        used_titles.add(safe_title.casefold())
        worksheet = workbook.create_sheet(safe_title)
        worksheet.append([f"{facility} roster · {target_month:%B %Y}"])
        worksheet.append(["Employee", "Activity", *[date(target_month.year, target_month.month, day).strftime("%d %a") for day in range(1, calendar.monthrange(target_month.year, target_month.month)[1] + 1)]])
        for row in assignment_rows:
            if row.get("sheet") != facility:
                continue
            worksheet.append([row.get("employee", ""), row.get("activity", ""), *row.get("assignments", [])])
        worksheet.freeze_panes = "C3"
        worksheet.auto_filter.ref = f"A2:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
        worksheet.column_dimensions["A"].width = 26
        worksheet.column_dimensions["B"].width = 20
        for column in range(3, worksheet.max_column + 1):
            worksheet.column_dimensions[get_column_letter(column)].width = 9

    summary = workbook.create_sheet("Validation Summary")
    summary.append(["Roster Validation Summary"])
    summary.append(["Target month", f"{target_month:%B %Y}"])
    summary.append(["Overall status", validation_report.get("overall", "NEEDS_ATTENTION")])
    summary.append(["Rule compliance", validation_report.get("rules", "NOT VERIFIED")])
    summary.append(["Staffing", validation_report.get("staffing", "NOT VERIFIED")])
    summary.append(["Historical completeness", validation_report.get("history", "NOT VERIFIED")])
    summary.append([])
    summary.append(["Facility", "Date", "Shift", "Required", "Assigned", "Shortage", "Excess"])
    for item in validation_report.get("coverage", []):
        summary.append([item.get("facility"), item.get("date"), item.get("shift"), item.get("required"), item.get("actual"), item.get("shortage"), item.get("excess")])
    summary.append([])
    summary.append(["Severity", "Code", "Finding"])
    for item in warnings:
        summary.append([str(item.get("severity", "info")).upper(), item.get("code", ""), item.get("message", "")])
    summary.freeze_panes = "A9"
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 25
    summary.column_dimensions["C"].width = 100
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output, {"updated_cells": 0, "layout_warnings": []}
